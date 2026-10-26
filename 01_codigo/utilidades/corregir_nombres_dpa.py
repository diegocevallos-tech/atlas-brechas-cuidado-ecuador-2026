"""
corregir_nombres_dpa.py
=======================
Reasigna el nombre del cantón en el Dataset Maestro a partir del código DPA
(clave de unión), usando una tabla DPA oficial. NO recalcula D1, D2, D3 ni IBTC.

INPUT
  - XLSX maestro (hoja de datos con columnas: código DPA, cantón, provincia, IBTC, ...)
  - CSV DPA oficial con columnas: cod, canton, provincia   (ver CONFIG.DPA_CSV)

OUTPUT (carpeta CONFIG.OUT_DIR)
  - <xlsx_original>.xlsx             : mismo formato, nombres corregidos + Top20 regenerado
  - <csv_final>.csv                  : CSV con código DPA (4 caracteres) + nombre de cantón y provincia
  - reporte_cruce_codigo_nombre.csv: una fila por cantón (nombre publicado vs corregido)

USO
  python corregir_nombres_dpa.py --root <raíz_del_repo> --dpa <ruta_dpa.csv> [--out <carpeta>]

QUÉ EDITAR: solo el bloque CONFIG. QUÉ NO: la lógica de las funciones.
"""

import argparse
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd

# ───────────────────────── CONFIG ─────────────────────────
CONFIG = {
    # Rutas relativas a --root
    "XLSX_REL": "02_datos/IBTC_Dataset_Maestro_DAT4CCION2026.xlsx",
    "CSV_REL": "02_datos/ibtc_cantones_final.csv",
    "OUT_DIR": "salida_corregida",
    # Hoja de datos y hoja Top N
    "HOJA_DATOS": "IBTC_221_Cantones",
    "FILA_ENCABEZADO": 3,          # fila (1-based) con los encabezados de la hoja de datos
    "HOJA_TOP": "Top20_Mayor_Brecha",
    "FILA_ENCABEZADO_TOP": 2,      # fila (1-based) con encabezados de la hoja Top
    "TOP_N": 20,
    # Nombres de columnas (se validan antes de procesar)
    "COL_COD": "Cód. DPA",
    "COL_CANTON": "Cantón",
    "COL_PROV": "Provincia",
    "COL_IBTC": "IBTC",
    "COL_D1": "D1\nDemanda",
    "COL_D2": "D2\nOferta",
    "COL_D3": "D3\nVulnerab.",
    "COL_TIPO": "Tipología",
    # CSV del repositorio: columna de código
    "CSV_COL_COD": "CANTON",
    # Columnas de la hoja Top (en el orden de la hoja)
    "TOP_COLS": ["Rango", "Cantón", "Provincia", "IBTC", "D1", "D2", "D3", "Tipología"],
    "TOP_DECIMALES": 3,
    # Códigos que pueden diferir entre la tabla DPA y el Censo 2022 (código_tabla: código_censo)
    # La Concordia: 0808 en capas antiguas; INEC Censo 2022 la clasifica como 2302.
    "RECODIFICAR_DPA": {"0808": "2302"},
    # Equivalencias de nomenclatura (nombre DPA normalizado -> nombre normalizado usado en el XLSX)
    "ALIAS_DPA_A_XLSX": {
        "ORELLANA": "FRANCISCO DE ORELLANA",
        "QUITO": "DISTRITO METROPOLITANO DE QUITO",
        "EMPALME": "EL EMPALME",
        "GNRAL ANTONIO ELIZALDE": "GENERAL ANTONIO ELIZALDE",
        "CRNEL MARCELINO MARIDUENA": "CORONEL MARCELINO MARIDUENA",
    },
    # Prefijos de código que no son cantones (zonas no delimitadas)
    "PREFIJOS_EXCLUIR": ("90",),
    "SUFIJO_XLSX": "",   # "" = mismo nombre que el original (salida en otra carpeta)
    "SUFIJO_CSV": "",    # "" = mismo nombre que el original
}
# ──────────────────────────────────────────────────────────


def norm(s: str) -> str:
    """Normaliza un nombre: mayúsculas, sin tildes, sin paréntesis ni signos."""
    s = unicodedata.normalize("NFKD", str(s).upper()).encode("ascii", "ignore").decode()
    s = re.sub(r"\(.*?\)", "", s)
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def exigir_columnas(df: pd.DataFrame, cols, origen: str):
    faltan = [c for c in cols if c not in df.columns]
    if faltan:
        raise KeyError(f"[{origen}] Columnas no encontradas: {faltan}. Disponibles: {list(df.columns)}")


def cargar_dpa(ruta: Path, cfg: dict) -> pd.DataFrame:
    if not ruta.exists():
        raise FileNotFoundError(f"Tabla DPA no encontrada: {ruta}")
    dpa = pd.read_csv(ruta, dtype={"cod": str})
    exigir_columnas(dpa, ["cod", "canton", "provincia"], "DPA")
    dpa["cod"] = dpa["cod"].str.zfill(4).replace(cfg["RECODIFICAR_DPA"])
    dpa = dpa[~dpa["cod"].str.startswith(cfg["PREFIJOS_EXCLUIR"])].copy()
    if dpa["cod"].duplicated().any():
        raise ValueError("Códigos DPA duplicados en la tabla oficial.")
    return dpa


def cargar_maestro(ruta: Path, cfg: dict) -> pd.DataFrame:
    if not ruta.exists():
        raise FileNotFoundError(f"XLSX no encontrado: {ruta}")
    df = pd.read_excel(ruta, sheet_name=cfg["HOJA_DATOS"], header=cfg["FILA_ENCABEZADO"] - 1,
                       dtype={cfg["COL_COD"]: str})
    exigir_columnas(df, [cfg["COL_COD"], cfg["COL_CANTON"], cfg["COL_PROV"], cfg["COL_IBTC"]], "XLSX")
    df = df[df[cfg["COL_COD"]].notna()].copy()
    df["_cod"] = df[cfg["COL_COD"]].str.zfill(4)
    return df


def construir_cruce(df: pd.DataFrame, dpa: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    """Devuelve una fila por código con el nombre publicado y el nombre corregido (con tildes del XLSX)."""
    solo_x = set(df["_cod"]) - set(dpa["cod"])
    solo_d = set(dpa["cod"]) - set(df["_cod"])
    if solo_x or solo_d:
        raise ValueError(f"Códigos que no cruzan. Solo en XLSX: {sorted(solo_x)} | solo en DPA: {sorted(solo_d)}")

    # Diccionario (provincia_2dig, nombre normalizado) -> nombre con tildes tal como figura en el XLSX.
    # Los nombres del XLSX son el conjunto correcto por provincia; solo estaban mal asignados.
    pool = {}
    for _, r in df.iterrows():
        pool[(r["_cod"][:2], norm(r[cfg["COL_CANTON"]]))] = r[cfg["COL_CANTON"]]

    filas = []
    for _, r in dpa.iterrows():
        n = norm(r["canton"])
        n = norm(cfg["ALIAS_DPA_A_XLSX"].get(n, n))
        clave = (r["cod"][:2], n)
        if clave not in pool:
            raise KeyError(f"Nombre DPA sin equivalente en el XLSX: {r['cod']} {r['canton']} (alias faltante en CONFIG)")
        filas.append({"_cod": r["cod"], "nombre_corregido": pool[clave], "nombre_dpa": r["canton"]})
    cruce = pd.DataFrame(filas)
    if cruce["nombre_corregido"].duplicated().sum() > 0:
        # Olmedo y Bolívar existen en dos provincias: duplicado legítimo
        dup = cruce[cruce["nombre_corregido"].duplicated(keep=False)]
        if not set(dup["_cod"].str[:2]).issuperset({"11", "13"}) and len(dup) > 4:
            raise ValueError(f"Duplicados inesperados de nombre:\n{dup}")

    out = df.merge(cruce, on="_cod", how="left")
    out["nombre_publicado"] = out[cfg["COL_CANTON"]]
    out["nombre_ok"] = out["nombre_publicado"].map(norm) == out["nombre_corregido"].map(norm)
    return out


def escribir_xlsx(df: pd.DataFrame, origen: Path, destino: Path, cfg: dict):
    wb = openpyxl.load_workbook(origen)
    ws = wb[cfg["HOJA_DATOS"]]
    enc = {str(c.value): c.column for c in ws[cfg["FILA_ENCABEZADO"]] if c.value}
    for k in (cfg["COL_COD"], cfg["COL_CANTON"]):
        if k not in enc:
            raise KeyError(f"Encabezado '{k}' no encontrado en {cfg['HOJA_DATOS']}: {list(enc)}")
    corr = dict(zip(df["_cod"], df["nombre_corregido"]))
    for fila in range(cfg["FILA_ENCABEZADO"] + 1, ws.max_row + 1):
        cod = ws.cell(fila, enc[cfg["COL_COD"]]).value
        if cod is None:
            continue
        cod = str(cod).zfill(4)
        if cod in corr:
            ws.cell(fila, enc[cfg["COL_CANTON"]]).value = corr[cod]

    # Regenerar hoja Top N desde los datos corregidos
    top = df.sort_values(cfg["COL_IBTC"], ascending=False).head(cfg["TOP_N"])
    wt = wb[cfg["HOJA_TOP"]]
    d = cfg["TOP_DECIMALES"]
    for i, (_, r) in enumerate(top.iterrows(), start=1):
        fila = cfg["FILA_ENCABEZADO_TOP"] + i
        valores = [i, r["nombre_corregido"], r[cfg["COL_PROV"]], round(r[cfg["COL_IBTC"]], d),
                   round(r[cfg["COL_D1"]], d), round(r[cfg["COL_D2"]], d), round(r[cfg["COL_D3"]], d),
                   r[cfg["COL_TIPO"]]]
        for j, v in enumerate(valores, start=1):
            wt.cell(fila, j).value = v
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def escribir_csv(csv_in: Path, df: pd.DataFrame, destino: Path, cfg: dict):
    if not csv_in.exists():
        print(f"[AVISO] CSV no encontrado, se omite: {csv_in}")
        return
    c = pd.read_csv(csv_in, dtype={cfg["CSV_COL_COD"]: str}, encoding="utf-8-sig")
    exigir_columnas(c, [cfg["CSV_COL_COD"]], "CSV")
    c["_cod"] = c[cfg["CSV_COL_COD"]].str.zfill(4)
    nombres = df[["_cod", "nombre_corregido", cfg["COL_PROV"]]].rename(
        columns={"nombre_corregido": "NOMBRE_CANTON", cfg["COL_PROV"]: "NOMBRE_PROVINCIA"})
    c = c.merge(nombres, on="_cod", how="left")
    c[cfg["CSV_COL_COD"]] = c["_cod"]          # código DPA como texto de 4 caracteres (con cero inicial)
    c = c.drop(columns="_cod")
    if c["NOMBRE_CANTON"].isna().any():
        raise ValueError("Hay códigos del CSV sin nombre corregido.")
    c.to_csv(destino, index=False, encoding="utf-8-sig")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", required=True, help="Raíz del repositorio del Atlas")
    ap.add_argument("--dpa", required=True, help="CSV DPA oficial (cod, canton, provincia)")
    ap.add_argument("--out", default=None, help="Carpeta de salida (por defecto CONFIG.OUT_DIR bajo --root)")
    a = ap.parse_args()
    cfg = CONFIG
    root = Path(a.root).resolve()
    out = Path(a.out).resolve() if a.out else root / cfg["OUT_DIR"]
    out.mkdir(parents=True, exist_ok=True)

    xlsx = root / cfg["XLSX_REL"]
    csv = root / cfg["CSV_REL"]
    dpa = cargar_dpa(Path(a.dpa), cfg)
    df = cargar_maestro(xlsx, cfg)
    df = construir_cruce(df, dpa, cfg)

    n, mal = len(df), int((~df["nombre_ok"]).sum())
    print(f"Cantones: {n} | nombres incorrectos: {mal} | correctos: {n - mal}")

    rep = df[["_cod", "nombre_publicado", "nombre_corregido", cfg["COL_PROV"], cfg["COL_IBTC"],
              cfg["COL_TIPO"], "nombre_ok"]].rename(columns={"_cod": "cod_dpa"})
    rep.to_csv(out / "reporte_cruce_codigo_nombre.csv", index=False, encoding="utf-8-sig")
    escribir_xlsx(df, xlsx, out / f"{xlsx.stem}{cfg['SUFIJO_XLSX']}{xlsx.suffix}", cfg)
    escribir_csv(csv, df, out / f"{csv.stem}{cfg['SUFIJO_CSV']}.csv", cfg)
    print(f"Salidas en: {out}")


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, KeyError, ValueError) as e:
        sys.exit(f"[ERROR] {e}")
